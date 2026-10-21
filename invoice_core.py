"""Generic invoice validation, numbering, and rendering.

This module contains no seller, customer, product, or bank master data.
"""

from __future__ import annotations

import io
import re
import sqlite3
from contextlib import closing
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

CENT = Decimal("0.01")
MAX_LINES = 50
MAX_TEXT = 240


def clean_text(value: object, label: str, max_length: int = MAX_TEXT) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{label} must be text")
    value = value.strip()
    if not value or len(value) > max_length or any(ord(ch) < 32 for ch in value):
        raise ValueError(f"{label} is empty, too long, or contains control characters")
    return value


def money(value: object) -> Decimal:
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError("Unit price must be a number") from None
    if not amount.is_finite() or amount <= 0 or amount > Decimal("100000000000"):
        raise ValueError("Unit price is out of range")
    if amount != amount.quantize(CENT):
        raise ValueError("Unit price must have at most two decimal places")
    return amount


def validate_invoice(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("Invoice must be a JSON object")
    customer = clean_text(payload.get("customer"), "Customer")
    try:
        invoice_date = date.fromisoformat(payload.get("date", ""))
    except (ValueError, TypeError):
        raise ValueError("Date must use YYYY-MM-DD") from None
    currency = payload.get("currency")
    if currency not in {"NGN", "USD"}:
        raise ValueError("Currency must be NGN or USD")
    vat_included = payload.get("vat_included", False)
    if not isinstance(vat_included, bool):
        raise ValueError("vat_included must be true or false")
    raw_lines = payload.get("lines")
    if not isinstance(raw_lines, list) or not 1 <= len(raw_lines) <= MAX_LINES:
        raise ValueError(f"Provide 1 to {MAX_LINES} line items")
    lines = []
    total = Decimal("0.00")
    for index, raw in enumerate(raw_lines, 1):
        if not isinstance(raw, dict):
            raise ValueError(f"Line {index} must be an object")
        description = clean_text(raw.get("description"), f"Line {index} description")
        qty = raw.get("quantity")
        if type(qty) is not int or not 1 <= qty <= 100000:
            raise ValueError(f"Line {index} quantity must be a positive integer")
        unit_price = money(raw.get("unit_price"))
        spec_id = raw.get("spec_id") or None
        if spec_id is not None and (not isinstance(spec_id, str) or not re.fullmatch(r"[0-9a-f]{32}", spec_id)):
            raise ValueError(f"Line {index} has an invalid spec ID")
        amount = unit_price * qty
        total += amount
        lines.append({"description": description, "quantity": qty, "unit_price": unit_price,
                      "amount": amount, "spec_id": spec_id})
    if total > Decimal("1000000000000"):
        raise ValueError("Invoice total is out of range")
    vat = (total - total / Decimal("1.075")).quantize(CENT, rounding=ROUND_HALF_UP) if vat_included else Decimal("0.00")
    return {"customer": customer, "date": invoice_date, "currency": currency,
            "vat_included": vat_included, "lines": lines, "total": total, "vat": vat,
            "net": total - vat}


def next_number(database: Path, invoice_date: date) -> str:
    database.parent.mkdir(parents=True, exist_ok=True)
    with closing(sqlite3.connect(database, timeout=15)) as conn:
        with conn:
            conn.execute("CREATE TABLE IF NOT EXISTS counters (year INTEGER PRIMARY KEY, next_value INTEGER NOT NULL)")
            conn.execute("BEGIN IMMEDIATE")
            row = conn.execute("SELECT next_value FROM counters WHERE year=?", (invoice_date.year,)).fetchone()
            number = row[0] if row else 1
            conn.execute("INSERT INTO counters(year, next_value) VALUES(?, ?) ON CONFLICT(year) DO UPDATE SET next_value=excluded.next_value",
                         (invoice_date.year, number + 1))
    return f"PFI-{invoice_date.year}-{number:04d}"


def render_xlsx(invoice: dict, number: str, seller: str, details: str, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Proforma Invoice"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:E1")
    ws["A1"] = "PROFORMA INVOICE"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A2:E2")
    ws["A2"] = seller
    ws["A3"] = "Seller details"
    ws.merge_cells("B3:E3")
    ws["B3"] = details
    ws["A5"], ws["B5"] = "Invoice no.", number
    ws["D5"], ws["E5"] = "Date", invoice["date"].isoformat()
    ws["A6"], ws["B6"] = "Customer", invoice["customer"]
    ws["D6"], ws["E6"] = "Currency", invoice["currency"]
    for col, label in enumerate(["Description", "Quantity", "Unit price", "Amount", "Spec PDF"], 1):
        cell = ws.cell(8, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315F86")
    for row, line in enumerate(invoice["lines"], 9):
        ws.cell(row, 1, line["description"])
        ws.cell(row, 2, line["quantity"])
        ws.cell(row, 3, float(line["unit_price"]))
        ws.cell(row, 4, float(line["amount"]))
        ws.cell(row, 5, "Attached" if line["spec_id"] else "")
        for col in (3, 4):
            ws.cell(row, col).number_format = '#,##0.00'
    end = 9 + len(invoice["lines"])
    for row, title, value in [(end, "Total (VAT included if selected)", invoice["total"]),
                              (end + 1, "VAT included (7.5%)", invoice["vat"]),
                              (end + 2, "Net before VAT", invoice["net"])]:
        ws.cell(row, 3, title)
        ws.cell(row, 4, float(value)).number_format = '#,##0.00'
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 33
    ws.column_dimensions["D"].width = 19
    ws.column_dimensions["E"].width = 22
    ws.freeze_panes = "A9"
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:E{end+2}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def render_pdf(invoice: dict, number: str, seller: str, details: str, path: Path,
               spec_dir: Path) -> None:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=(595.28, 841.89), rightMargin=36, leftMargin=36,
                            topMargin=40, bottomMargin=40)
    visible_text = seller + details + invoice["customer"] + "".join(line["description"] for line in invoice["lines"])
    font_name = "Helvetica"
    if any(ord(character) > 255 for character in visible_text):
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        font_name = "STSong-Light"
    normal = ParagraphStyle("normal", fontName=font_name, fontSize=9, leading=13)
    right = ParagraphStyle("right", parent=normal, alignment=TA_RIGHT)
    title = ParagraphStyle("title", parent=normal, fontSize=17, leading=23, spaceAfter=12)
    def p(text: object, style=normal):
        from xml.sax.saxutils import escape
        return Paragraph(escape(str(text)), style)
    content = [p("PROFORMA INVOICE", title), p(seller), p(details), Spacer(1, 14),
               p(f"Invoice no.: {number}"), p(f"Date: {invoice['date'].isoformat()}"),
               p(f"Customer: {invoice['customer']}"), p(f"Currency: {invoice['currency']}"), Spacer(1, 14)]
    rows = [[p("Description"), p("Qty"), p("Unit price", right), p("Amount", right)]]
    for line in invoice["lines"]:
        rows.append([p(line["description"]), p(line["quantity"]),
                     p(f"{line['unit_price']:,.2f}", right), p(f"{line['amount']:,.2f}", right)])
    table = Table(rows, colWidths=[294, 42, 91, 96], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE8F2")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.HexColor("#315F86")),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FB")]),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    content += [table, Spacer(1, 13), p(f"Total: {invoice['currency']} {invoice['total']:,.2f}", right),
                p(f"VAT included (7.5%): {invoice['currency']} {invoice['vat']:,.2f}", right),
                p(f"Net before VAT: {invoice['currency']} {invoice['net']:,.2f}", right)]
    doc.build(content)
    buffer.seek(0)
    writer = PdfWriter()
    writer.append(PdfReader(buffer))
    seen = set()
    for line in invoice["lines"]:
        spec_id = line["spec_id"]
        if spec_id and spec_id not in seen:
            spec_path = spec_dir / f"{spec_id}.pdf"
            if not spec_path.is_file():
                raise ValueError("A selected spec PDF is missing")
            writer.append(str(spec_path))
            seen.add(spec_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("wb") as target:
        writer.write(target)

