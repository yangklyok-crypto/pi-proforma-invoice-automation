import io
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

from invoice_core import next_number, render_pdf, render_xlsx, validate_invoice


SAMPLE = Path(__file__).resolve().parents[1] / "examples" / "sample_invoice.json"


class InvoiceTests(unittest.TestCase):
    def test_inclusive_vat_and_files(self):
        payload = json.loads(SAMPLE.read_text(encoding="utf-8"))
        invoice = validate_invoice(payload)
        self.assertEqual(str(invoice["total"]), "2250.00")
        self.assertEqual(str(invoice["vat"]), "156.98")
        self.assertEqual(str(invoice["net"]), "2093.02")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            number = next_number(root / "numbers.sqlite3", invoice["date"])
            self.assertEqual(number, "PFI-2026-0001")
            self.assertEqual(next_number(root / "numbers.sqlite3", date(2026, 2, 1)), "PFI-2026-0002")
            pdf = root / "invoice.pdf"
            xlsx = root / "invoice.xlsx"
            render_pdf(invoice, number, "Example Supplier", "", pdf, root)
            render_xlsx(invoice, number, "Example Supplier", "", xlsx)
            self.assertEqual(len(PdfReader(str(pdf)).pages), 1)
            ws = load_workbook(xlsx, data_only=True).active
            self.assertEqual(ws["B5"].value, number)
            self.assertEqual(ws["D9"].value, 2000)

    def test_spec_append_once_for_repeated_item(self):
        payload = json.loads(SAMPLE.read_text(encoding="utf-8"))
        payload["vat_included"] = False
        payload["lines"][0]["spec_id"] = "a" * 32
        payload["lines"][1]["spec_id"] = "a" * 32
        invoice = validate_invoice(payload)
        self.assertEqual(invoice["vat"], 0)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            writer = PdfWriter()
            writer.add_blank_page(width=500, height=700)
            with (root / ("a" * 32 + ".pdf")).open("wb") as output:
                writer.write(output)
            target = root / "invoice.pdf"
            render_pdf(invoice, "PFI-2026-0001", "Example Supplier", "", target, root)
            self.assertEqual(len(PdfReader(str(target)).pages), 2)

    def test_unicode_invoice(self):
        payload = json.loads(SAMPLE.read_text(encoding="utf-8"))
        payload["customer"] = "示例买家"
        payload["lines"][0]["description"] = "示例设备"
        invoice = validate_invoice(payload)
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "invoice.pdf"
            render_pdf(invoice, "PFI-2026-0001", "示例供应商", "", target, Path(directory))
            self.assertEqual(len(PdfReader(str(target)).pages), 1)

    def test_bad_values(self):
        payload = json.loads(SAMPLE.read_text(encoding="utf-8"))
        for value in ("NaN", "-1", "0.001"):
            altered = json.loads(json.dumps(payload))
            altered["lines"][0]["unit_price"] = value
            with self.assertRaises(ValueError):
                validate_invoice(altered)
        payload["lines"][0]["spec_id"] = "../../confidential.pdf"
        with self.assertRaises(ValueError):
            validate_invoice(payload)


if __name__ == "__main__":
    unittest.main()

